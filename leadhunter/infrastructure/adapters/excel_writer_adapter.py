"""Excel writer adapter — Infrastructure layer (REQ-056-062).

Writes lead records to an Excel file with a soft, smooth 5-column pastel format,
without harsh cell borders, matching user's exact screenshot (Ảnh 1).
"""

from __future__ import annotations

import logging
import os
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from leadhunter.domain.entities.lead import Lead

from leadhunter.domain.exceptions import FileWriteError

logger = logging.getLogger(__name__)


class ExcelWriterAdapter:
    """Writes lead records to a soft, smooth Excel (.xlsx) file with seamless pastel column fills."""

    def write(self, leads: list["Lead"], output_path: str) -> None:
        """Write leads to Excel workbook.

        Args:
            leads: List of Lead entities.
            output_path: Target file path (must end in .xlsx).
        """
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill

        out_path = Path(output_path)
        tmp_path = out_path.with_suffix(".tmp")

        try:
            out_path.parent.mkdir(parents=True, exist_ok=True)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Danh Sách Lead"

            # Column headers: SĐT | Tên Cửa Hàng | Địa chỉ | Link Google Maps | Tình trạng
            headers = [
                "Số Điện Thoại",
                "Tên Cửa Hàng",
                "Địa Chỉ",
                "Link Google Maps",
                "Tình Trạng",
            ]
            ws.append(headers)

            # Header styling: Navy Blue background with Bold White text
            header_fill = PatternFill(
                start_color="1F497D", end_color="1F497D", fill_type="solid"
            )
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_align = Alignment(horizontal="left", vertical="center")

            # Apply header styling
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align

            # Data rows
            for lead in leads:
                phone_str = (
                    lead.phone.value
                    if hasattr(lead.phone, "value")
                    else str(lead.phone or "")
                )
                if phone_str.startswith("+84"):
                    phone_str = "0" + phone_str[3:]

                link_src = ""
                if lead.website and hasattr(lead.website, "url") and lead.website.url and lead.website.url.startswith("http"):
                    link_src = lead.website.url
                elif lead.source_reference and lead.source_reference.startswith("http"):
                    link_src = lead.source_reference
                elif lead.source_reference and not lead.source_reference.startswith("/") and not lead.source_reference.startswith("file:") and not lead.source_reference.endswith(".xlsx") and not lead.source_reference.endswith(".csv"):
                    link_src = lead.source_reference
                else:
                    if lead.company_name:
                        import urllib.parse
                        q = urllib.parse.quote(f"{lead.company_name} {lead.address or ''}".strip())
                        link_src = f"https://www.google.com/maps/search/?api=1&query={q}"
                    else:
                        link_src = "https://www.google.com/maps"
                if lead.notes:
                    status_str = lead.notes
                elif (
                    lead.status
                    and hasattr(lead.status, "value")
                    and lead.status.value != "NEW"
                ):
                    status_str = lead.status.value
                else:
                    status_str = ""

                row = [
                    phone_str,
                    lead.company_name or "",
                    lead.address or "",
                    link_src,
                    status_str,
                ]
                ws.append(row)

            # Header row height: 24pt
            ws.row_dimensions[1].height = 24

            # Sleek Column Widths matching Ảnh 1
            col_widths = {"A": 20, "B": 45, "C": 45, "D": 18, "E": 35}
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width

            # 5-Column Soft Pastel Fills matching Ảnh 1
            fill_colors = {
                1: PatternFill(
                    start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"
                ),  # SĐT: Soft Green
                2: PatternFill(
                    start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
                ),  # Tên: Soft Yellow
                3: PatternFill(
                    start_color="C9DAF8", end_color="C9DAF8", fill_type="solid"
                ),  # Địa chỉ: Soft Blue
                4: PatternFill(
                    start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"
                ),  # Link: Soft Red/Pink
                5: PatternFill(
                    start_color="D9D2E9", end_color="D9D2E9", fill_type="solid"
                ),  # Tình trạng: Soft Purple
            }

            # Regular Calibri 11pt font
            data_font = Font(name="Calibri", size=11, color="000000")
            # Single-line display (wrap_text=False) for smooth rows
            data_align = Alignment(
                horizontal="left", vertical="center", wrap_text=False
            )

            # Compact 20pt Row Height
            for row_idx in range(2, ws.max_row + 1):
                ws.row_dimensions[row_idx].height = 20

            # Apply smooth pastel fills without harsh cell borders
            for row in ws.iter_rows(
                min_row=2, max_row=ws.max_row, min_col=1, max_col=5
            ):
                for cell in row:
                    cell.font = data_font
                    cell.alignment = data_align
                    if cell.column in fill_colors:
                        cell.fill = fill_colors[cell.column]

            # Freeze top row
            ws.freeze_panes = "A2"

            # Enable standard gridlines view
            if ws.views.sheetView:
                ws.views.sheetView[0].showGridLines = True

            # Write to temp file first (atomic write)
            wb.save(str(tmp_path))
            os.replace(str(tmp_path), str(out_path))

            logger.info(
                "Excel file written successfully",
                extra={
                    "context": {
                        "output_path": str(out_path),
                        "record_count": len(leads),
                    }
                },
            )

        except Exception as exc:
            if tmp_path.exists():
                try:
                    tmp_path.unlink()
                except OSError:
                    pass
            raise FileWriteError(str(out_path), exc) from exc
