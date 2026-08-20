"""Word writer adapter — Infrastructure layer.

Writes Google Maps lead records to a .docx Word document formatted as a spacious 3-column A4 Landscape grid card layout.
Matches user requirement:
1. Standard readable font sizes (11.5pt Name, 10pt Address, 12.5pt Phone)
2. Expanded card box height & width with generous top/bottom/left/right padding (w:tcMar)
"""

from __future__ import annotations

import logging
import os
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from leadhunter.domain.entities.lead import Lead

from leadhunter.domain.exceptions import FileWriteError

logger = logging.getLogger(__name__)


class WordWriterAdapter:
    """Writes lead records to a .docx Word document with tall, spacious 3-column Cards."""

    def write(
        self,
        leads: list["Lead"],
        output_path: str,
        title: str = "DANH SÁCH LEAD GOOGLE MAPS",
    ) -> None:
        """Write leads to a Word (.docx) document with tall, spacious 3-column grid card layout."""
        import docx
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        out_path = Path(output_path)
        tmp_path = out_path.with_suffix(".tmp")

        try:
            out_path.parent.mkdir(parents=True, exist_ok=True)

            doc = docx.Document()

            # Set section to A4 Landscape with 0.4 inch narrow margins
            section = doc.sections[0]
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = Inches(11.69)
            section.page_height = Inches(8.27)
            section.top_margin = Inches(0.4)
            section.bottom_margin = Inches(0.4)
            section.left_margin = Inches(0.4)
            section.right_margin = Inches(0.4)

            # Top Header Box matching user request style
            today_str = datetime.now().strftime("%d/%m/%Y")
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_title.paragraph_format.space_after = Pt(8)
            run_title = p_title.add_run(f"📋 {title}")
            run_title.font.name = "Times New Roman"
            run_title.font.size = Pt(13.5)
            run_title.font.bold = True

            # Create 3-column Grid Table
            num_cols = 3
            num_leads = len(leads)
            num_rows = (num_leads + num_cols - 1) // num_cols if num_leads > 0 else 1

            table = doc.add_table(rows=num_rows, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            # 3 cols x 3.63 in = 10.89 in across A4 Landscape (100% full printable width)
            col_width = Inches(3.63)

            for r_idx in range(num_rows):
                for c_idx in range(num_cols):
                    lead_idx = r_idx * num_cols + c_idx
                    cell = table.cell(r_idx, c_idx)
                    cell.width = col_width

                    tcPr = cell._element.get_or_add_tcPr()

                    # Set cell border for crisp card boxes
                    tcBorders = tcPr.first_child_found_in("w:tcBorders")
                    if tcBorders is None:
                        tcBorders = OxmlElement("w:tcBorders")
                        tcPr.append(tcBorders)
                    for edge in ("top", "left", "bottom", "right"):
                        b_elem = OxmlElement(f"w:{edge}")
                        b_elem.set(qn("w:val"), "single")
                        b_elem.set(qn("w:sz"), "8")  # 1pt border
                        b_elem.set(qn("w:space"), "0")
                        b_elem.set(qn("w:color"), "777777")
                        tcBorders.append(b_elem)

                    # Set expanded cell margins (padding) to increase card box height and width
                    tcMar = OxmlElement("w:tcMar")
                    for m_name, dxa_val in (
                        ("top", "240"),
                        ("bottom", "240"),
                        ("left", "180"),
                        ("right", "180"),
                    ):
                        m_elem = OxmlElement(f"w:{m_name}")
                        m_elem.set(qn("w:w"), str(dxa_val))
                        m_elem.set(qn("w:type"), "dxa")
                        tcMar.append(m_elem)
                    tcPr.append(tcMar)

                    if lead_idx < num_leads:
                        lead = leads[lead_idx]

                        phone_str = (
                            lead.phone.value
                            if hasattr(lead.phone, "value")
                            else str(lead.phone or "")
                        )
                        if phone_str.startswith("+84"):
                            phone_str = "0" + phone_str[3:]

                        comp_name = (
                            lead.company_name or "CỬA HÀNG / SPA / CƠ SỞ"
                        ).upper()
                        addr = lead.address or "Chưa cập nhật địa chỉ"

                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(6)
                        p.paragraph_format.space_after = Pt(6)
                        p.paragraph_format.line_spacing = 1.3

                        # 1. Tên Cửa Hàng / Cơ Sở (Standard 11.5pt Bold Navy)
                        r_name = p.add_run(comp_name + "\n")
                        r_name.font.name = "Times New Roman"
                        r_name.font.size = Pt(11.5)
                        r_name.font.bold = True
                        r_name.font.color.rgb = RGBColor(15, 34, 64)

                        # 2. Địa chỉ (Standard 10pt)
                        r_addr_lbl = p.add_run("📍 ")
                        r_addr_lbl.font.name = "Times New Roman"
                        r_addr_lbl.font.size = Pt(10)

                        r_addr = p.add_run(addr + "\n")
                        r_addr.font.name = "Times New Roman"
                        r_addr.font.size = Pt(10)

                        # 3. Số Điện Thoại (SĐT Standard 12.5pt BOLD Pitch Black)
                        r_ph_lbl = p.add_run("📞 SĐT: ")
                        r_ph_lbl.font.name = "Times New Roman"
                        r_ph_lbl.font.size = Pt(11)
                        r_ph_lbl.font.bold = True

                        r_ph = p.add_run(phone_str + "\n")
                        r_ph.font.name = "Times New Roman"
                        r_ph.font.size = Pt(12.5)
                        r_ph.font.bold = True
                        r_ph.font.color.rgb = RGBColor(0, 0, 0)

                        # 4. Dòng Ghi Chú cho Telesale (9.5pt Italic)
                        note_text = (
                            f"📝 Note: {lead.notes}"
                            if lead.notes
                            else "📝 Ghi chú: ......................................."
                        )
                        r_note = p.add_run(note_text)
                        r_note.font.name = "Times New Roman"
                        r_note.font.size = Pt(9.5)
                        r_note.font.italic = True
                        r_note.font.color.rgb = RGBColor(100, 100, 100)

            doc.save(str(tmp_path))
            os.replace(str(tmp_path), str(out_path))

            logger.info(f"Word file written successfully to {out_path}")

        except Exception as exc:
            if tmp_path.exists():
                try:
                    tmp_path.unlink()
                except OSError:
                    pass
            raise FileWriteError(str(out_path), exc) from exc
