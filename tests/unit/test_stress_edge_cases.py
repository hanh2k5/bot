import pytest
from pathlib import Path
from leadhunter.domain.services.telecom_service import is_viettel, is_tong_dai, is_vinaphone, is_mobifone
from leadhunter.domain.services.normalization_service import (
    normalize_phone,
    normalize_company_name,
    normalize_address,
    normalize_website,
    normalize_contact_name,
)
from leadhunter.domain.exceptions import InvalidCompanyNameError
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from leadhunter.application.use_cases.auto_run_use_case import (
    _is_valid_phone,
    _has_website,
    _is_valid_name,
    _is_duplicate,
    _build_lead,
)
from leadhunter.infrastructure.adapters.excel_writer_adapter import ExcelWriterAdapter
from leadhunter.domain.entities.lead import Lead, LeadStatus
from gui_app import _is_node_junk


class TestPhoneNormalizationStressEdgeCases:
    """Stress tests & edge cases cho chuẩn hóa số điện thoại."""

    @pytest.mark.parametrize(
        "invalid_input",
        [
            None,
            "",
            "   ",
            "abc",
            "090123",  # Quá ngắn
            "0901234567890",  # Quá dài
            "1234567890",  # Không bắt đầu bằng 0
            "0000000000",  # Đầu số không tồn tại
            "+84(0)901234567",  # Cú pháp sai +84(0)
            "090-123-4567ext12",
        ],
    )
    def test_normalize_phone_handles_all_invalid_edge_cases(self, invalid_input):
        result = normalize_phone(invalid_input)
        assert result is None or result.value is None or result.value == ""

    @pytest.mark.parametrize(
        "raw_phone, expected_normalized",
        [
            ("0901234567", "0901234567"),
            ("+84901234567", "0901234567"),
            ("84901234567", "0901234567"),
            ("090.123.4567", "0901234567"),
            (" 090 123 4567 ", "0901234567"),
            ("(090) 123-4567", "0901234567"),
        ],
    )
    def test_normalize_phone_valid_variations(self, raw_phone, expected_normalized):
        vo = normalize_phone(raw_phone)
        assert vo is not None
        assert vo.value == expected_normalized


class TestCompanyNameNormalizationStressEdgeCases:
    """Stress tests & edge cases cho tên công ty/cửa hàng."""

    def test_company_name_empty_or_none_raises_or_returns_empty(self):
        with pytest.raises(InvalidCompanyNameError):
            normalize_company_name("")

        with pytest.raises(InvalidCompanyNameError):
            normalize_company_name("   ")

    def test_company_name_collapses_multiple_spaces(self):
        vo = normalize_company_name("  Công   Ty   TNHH   Xây   Dựng  ")
        assert "công ty tnhh xây dựng" in vo.value.lower()

    def test_is_valid_name_filters_junk_and_place_ids(self):
        assert _is_valid_name("0x89123abcdef") is False
        assert _is_valid_name("ChI123abcdef456") is False
        assert _is_valid_name("See nearby") is False
        assert _is_valid_name("Directions") is False
        assert _is_valid_name("Nha Khoa Răng Sứ") is True


class TestWebsiteAndAddressEdgeCases:
    """Stress tests cho website & địa chỉ null/corrupt data."""

    def test_has_website_robust_against_all_types(self):
        assert _has_website({"website": None}) is False
        assert _has_website({"website": 12345}) is False
        assert _has_website({"website": []}) is False
        assert _has_website({"website": {}}) is False
        assert _has_website({"website": "  "}) is False
        assert _has_website({"website": "https://example.com"}) is True

    def test_normalize_address_handles_none_and_long_strings(self):
        assert normalize_address(None) == ""
        assert normalize_address("   ") == ""
        long_addr = "123 " * 100
        norm = normalize_address(long_addr)
        assert len(norm) > 0
        assert "  " not in norm


class TestLeadEntityConstructionStress:
    """Stress tests cho hàm _build_lead khi nhận dữ liệu rác/thiếu từ Google Maps."""

    def test_build_lead_handles_missing_fields_gracefully(self):
        raw_bad = {
            "company_name": "",
            "phone": "0901234567",
            "address": "123 Quận 1",
        }
        lead = _build_lead(raw_bad, None, "batch_1")
        assert lead is None  # Invalid name or phone VO None returns None

    def test_build_lead_preserves_valid_fields(self):
        phone_vo = normalize_phone("0908238968")
        raw_good = {
            "company_name": "Nha Khoa Kim",
            "phone": "0908238968",
            "address": "33 Đường 3/2, Q10, TP.HCM",
            "website": "https://nhakhoakim.com",
            "source_reference": "https://maps.google.com/?cid=999",
        }
        lead = _build_lead(raw_good, phone_vo, "batch_1")
        assert lead is not None
        assert lead.company_name == "Nha Khoa Kim"
        assert lead.phone == "0908238968"
        assert lead.website == "https://nhakhoakim.com"
        assert lead.address == "33 Đường 3/2, Q10, TP.HCM"


class TestExcelWriterStressEdgeCases:
    """Stress tests cho bộ xuất file Excel khi gặp list rỗng hoặc ký tự đặc biệt."""

    def test_excel_writer_handles_empty_lead_list(self, tmp_path):
        writer = ExcelWriterAdapter()
        out_file = tmp_path / "empty_export.xlsx"
        writer.write([], str(out_file))

        assert out_file.exists()
        import openpyxl
        wb = openpyxl.load_workbook(out_file)
        sheet = wb.active
        # Header line exists
        assert sheet.max_row == 1
        headers = [cell.value for cell in sheet[1]]
        assert headers == ["Số Điện Thoại", "Tên Cửa Hàng", "Địa Chỉ", "Link Google Maps", "Tình Trạng"]

    def test_excel_writer_handles_special_characters(self, tmp_path):
        writer = ExcelWriterAdapter()
        lead_special = Lead(
            company_name="Cty Xây Dựng <Test> & \"Đại Việt\" 'Quận 7'",
            contact_name="",
            email="",
            phone="0918234567",
            website="",
            address="456/78 Đường 9A, Bình Hưng, Bình Chánh",
            source="google_maps",
            source_reference="https://maps.google.com/?cid=special&test=1",
            status=LeadStatus.NEW,
            import_batch_id="batch_special",
        )
        out_file = tmp_path / "special_export.xlsx"
        writer.write([lead_special], str(out_file))

        assert out_file.exists()
        import openpyxl
        wb = openpyxl.load_workbook(out_file)
        sheet = wb.active
        assert sheet.max_row == 2
        assert sheet.cell(row=2, column=2).value == "Cty Xây Dựng <Test> & \"Đại Việt\" 'Quận 7'"
