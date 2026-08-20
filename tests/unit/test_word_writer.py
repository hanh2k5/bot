"""Unit tests for WordWriterAdapter (REQ-WordExport)."""

import os
from pathlib import Path
import docx
from leadhunter.domain.entities.lead import Lead
from leadhunter.domain.value_objects.phone_number import PhoneNumber
from leadhunter.infrastructure.adapters.word_writer_adapter import WordWriterAdapter


def test_word_writer_creates_valid_docx(tmp_path: Path):
    adapter = WordWriterAdapter()
    leads = [
        Lead(
            phone=PhoneNumber("0983325969"),
            company_name="HỘ KINH DOANH TEST STUDIO",
            address="Số 123 Đường ABC, Q1, TP.HCM",
            source_reference="https://maps.google.com/test",
            contact_name="PHÙNG NGỌC HẢI",
            email="",
            website="",
            source="GMap",
        )
    ]
    out_file = tmp_path / "test_output.docx"
    adapter.write(leads, str(out_file))

    assert out_file.exists()
    assert out_file.stat().st_size > 0

    doc = docx.Document(str(out_file))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    cell_text = table.cell(0, 0).text
    assert "0983325969" in cell_text
    assert "HỘ KINH DOANH TEST STUDIO" in cell_text
